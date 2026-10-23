import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from . import tax
from .allocations import AllocationError, compute
from .importer import HEADERS
from .models import (AuditEvent, Category, Expense, GSTType, LegalEntity, SourceType, Status, SyncStatus)

User = get_user_model()
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


class TaxTests(TestCase):
    def test_srs_example_intra_and_inter_state(self):
        intra = tax.calculate(GSTType.CGST_SGST, Decimal("10000"), Decimal("18"))
        self.assertEqual(intra["cgst_amount"], Decimal("900.00"))
        self.assertEqual(intra["sgst_utgst_amount"], Decimal("900.00"))
        self.assertEqual(intra["gross_amount"], Decimal("11800.00"))
        inter = tax.calculate(GSTType.IGST, Decimal("10000"), Decimal("18"))
        self.assertEqual(inter["igst_amount"], Decimal("1800.00"))
        self.assertEqual(inter["gross_amount"], Decimal("11800.00"))

    def test_exempt_has_no_tax_and_rcm_not_added_to_gross(self):
        self.assertEqual(tax.calculate(GSTType.EXEMPT, Decimal("500"), Decimal("18"))["gross_amount"], Decimal("500.00"))
        rcm = tax.calculate(GSTType.RCM, Decimal("1000"), Decimal("18"), supplier_state="33", place_of_supply="33")
        self.assertEqual(rcm["gross_amount"], Decimal("1000.00"))
        self.assertEqual(rcm["rcm_liability"], Decimal("180.00"))

    def test_state_inference_is_only_a_hint(self):
        self.assertEqual(tax.infer_gst_type("33", "33"), GSTType.CGST_SGST)
        self.assertEqual(tax.infer_gst_type("33", "29"), GSTType.IGST)
        self.assertEqual(tax.infer_gst_type("", "29"), "")


class AllocationTests(TestCase):
    def test_percent_residual_goes_to_last_line(self):
        lines, rest = compute(Decimal("100.00"), [{"method": "PERCENT", "percent": Decimal("33.3333")},
                                                  {"method": "PERCENT", "percent": Decimal("33.3333")},
                                                  {"method": "PERCENT", "percent": Decimal("33.3334")}])
        self.assertEqual(sum(l["amount"] for l in lines), Decimal("100.00"))
        self.assertEqual(rest, Decimal("0"))

    def test_over_allocation_rejected(self):
        with self.assertRaises(AllocationError):
            compute(Decimal("100"), [{"method": "PERCENT", "percent": Decimal("70")},
                                     {"method": "PERCENT", "percent": Decimal("40")}])
        with self.assertRaises(AllocationError):
            compute(Decimal("100"), [{"method": "FIXED", "amount": Decimal("101")}])


class BaseAPITest(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", "--no-expenses", stdout=io.StringIO())
        cls.alpha = LegalEntity.objects.get(code="ALPHA")
        cls.beta = LegalEntity.objects.get(code="BETA")
        cls.meals = Category.objects.get(code="MEALS")
        cls.fuel = Category.objects.get(code="FUEL")  # receipt not required

    def login(self, username):
        self.client.force_login(User.objects.get(username=username))

    def create(self, **extra):
        body = {"title": "Team lunch", "expense_date": str(timezone.localdate() - timedelta(days=2)),
                "category": self.fuel.id, "paying_entity": self.alpha.id, "gross_amount": "1200.00"}
        body.update(extra)
        r = self.client.post("/api/expenses/", body, content_type="application/json")
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()

    def act(self, exp_id, action, comment=""):
        return self.client.post(f"/api/expenses/{exp_id}/transition/", {"action": action, "comment": comment},
                                content_type="application/json")


class ExpenseAPITests(BaseAPITest):
    def test_requires_login(self):
        self.assertIn(self.client.get("/api/expenses/").status_code, (401, 403))
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_create_with_gst_calculates_components(self):
        self.login("employee")
        e = self.create(gst_type="CGST_SGST", gst_rate="18", taxable_value="10000", gross_amount=None)
        self.assertEqual(e["cgst_amount"], "900.00")
        self.assertEqual(e["gross_amount"], "11800.00")
        self.assertFalse(e["tax_variance"])
        self.assertEqual(e["verified_itc_amount"], "0.00")  # never inferred
        self.assertTrue(e["ref"].startswith("EXP-"))

    def test_declared_tax_mismatch_flags_variance(self):
        self.login("employee")
        e = self.create(gst_type="IGST", gst_rate="18", taxable_value="10000", igst_amount="2000", gross_amount="12000")
        self.assertTrue(e["tax_variance"])
        self.assertEqual(e["source_tax_values"]["igst_amount"], "2000.00")
        self.assertEqual(e["calculated_tax_values"]["igst_amount"], "1800.00")

    def test_full_workflow_and_lock(self):
        self.login("employee")
        e = self.create()
        self.assertEqual(self.act(e["id"], "approve").status_code, 400)  # cannot approve own / wrong state
        self.assertEqual(self.act(e["id"], "submit").status_code, 200)
        self.login("employee")
        self.assertEqual(self.client.patch(f"/api/expenses/{e['id']}/", {"title": "x"},
                                           content_type="application/json").status_code, 403)
        self.login("viewer")
        self.assertEqual(self.act(e["id"], "approve").status_code, 400)
        self.login("manager")
        r = self.act(e["id"], "approve", "ok")
        self.assertEqual(r.json()["status"], Status.MANAGER_APPROVED)
        self.login("accountant")
        self.assertEqual(self.act(e["id"], "verify").json()["status"], Status.FINANCE_VERIFIED)
        self.login("finance")
        self.assertEqual(self.client.patch(f"/api/expenses/{e['id']}/", {"title": "x"},
                                           content_type="application/json").status_code, 403)
        self.assertEqual(self.act(e["id"], "reopen").status_code, 400)  # reason required
        self.assertEqual(self.act(e["id"], "reopen", "Wrong category").json()["status"], Status.RETURNED)
        exp = Expense.objects.get(pk=e["id"])
        self.assertEqual([w.action for w in exp.workflow_events.all()], ["submit", "approve", "verify", "reopen"])
        self.assertTrue(AuditEvent.objects.filter(action="expense.verify", record_id=e["id"]).exists())

    def test_receipt_required_before_submit(self):
        self.login("employee")
        e = self.create(category=self.meals.id, gross_amount="800")
        r = self.act(e["id"], "submit")
        self.assertEqual(r.status_code, 400)
        self.assertIn("receipt", r.json()["detail"].lower())
        up = self.client.post(f"/api/expenses/{e['id']}/attachments/",
                              {"file": SimpleUploadedFile("r.pdf", PDF, content_type="application/pdf")})
        self.assertEqual(up.status_code, 201, up.content)
        self.assertEqual(self.act(e["id"], "submit").status_code, 200)

    def test_attachment_type_is_sniffed_not_trusted(self):
        self.login("employee")
        e = self.create()
        r = self.client.post(f"/api/expenses/{e['id']}/attachments/",
                             {"file": SimpleUploadedFile("evil.pdf", b"<script>", content_type="application/pdf")})
        self.assertEqual(r.status_code, 400)

    def test_manager_approval_limit(self):
        self.login("employee2")  # BETA, manager limit 50,000 there
        cat = Category.objects.create(code="EQUIP", name="Equipment", receipt_required=False)
        e = self.client.post("/api/expenses/", {"title": "Big", "expense_date": str(timezone.localdate()),
                                                "category": cat.id, "paying_entity": self.beta.id,
                                                "gross_amount": "60000"}, content_type="application/json").json()
        self.assertEqual(self.act(e["id"], "submit").status_code, 200)
        self.login("manager")
        r = self.act(e["id"], "approve")
        self.assertEqual(r.status_code, 400)
        self.assertIn("limit", r.json()["detail"])

    def test_entity_scope(self):
        self.login("employee")
        e = self.create()
        self.login("employee2")  # BETA only
        self.assertEqual(self.client.get(f"/api/expenses/{e['id']}/").status_code, 404)
        self.assertEqual(self.client.get("/api/expenses/").json()["count"], 0)
        r = self.client.post("/api/expenses/", {"title": "x", "expense_date": "2026-01-01", "category": self.fuel.id,
                                                "paying_entity": self.alpha.id, "gross_amount": "10"},
                             content_type="application/json")
        self.assertEqual(r.status_code, 400)
        self.login("auditor")
        self.assertEqual(self.client.get(f"/api/expenses/{e['id']}/").status_code, 200)
        self.assertEqual(self.client.post("/api/expenses/", {}, content_type="application/json").status_code, 403)

    def test_itc_review_controls(self):
        self.login("employee")
        e = self.create(gst_type="CGST_SGST", gst_rate="18", taxable_value="1000", gross_amount=None)
        self.act(e["id"], "submit")
        r = self.client.post(f"/api/expenses/{e['id']}/itc-review/", {"itc_status": "VERIFIED_ELIGIBLE",
                             "verified_itc_amount": "180", "reason": "x"}, content_type="application/json")
        self.assertEqual(r.status_code, 403)  # submitter cannot review
        self.login("accountant")
        r = self.client.post(f"/api/expenses/{e['id']}/itc-review/", {"itc_status": "VERIFIED_ELIGIBLE",
                             "verified_itc_amount": "500", "reason": "x"}, content_type="application/json")
        self.assertEqual(r.status_code, 400)  # exceeds GST
        r = self.client.post(f"/api/expenses/{e['id']}/itc-review/", {"itc_status": "VERIFIED_ELIGIBLE",
                             "verified_itc_amount": "180", "reason": "Valid invoice", "reference": "GSTR2B Oct"},
                             content_type="application/json")
        self.assertEqual(r.json()["verified_itc_amount"], "180.00")

    def test_allocation_api(self):
        self.login("employee")
        e = self.create(gross_amount="1000")
        r = self.client.put(f"/api/expenses/{e['id']}/allocations/", {"lines": [
            {"beneficiary_entity": self.alpha.id, "method": "PERCENT", "percent": "60"},
            {"beneficiary_entity": self.beta.id, "method": "PERCENT", "percent": "40"}]},
            content_type="application/json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual([a["amount"] for a in r.json()["allocations"]], ["600.00", "400.00"])
        self.login("finance")
        d = self.client.get("/api/dashboard/").json()
        self.assertEqual(d["kpis"]["gross"], "1000.00")  # invoice counted once
        alloc = {x["label"]: x["value"] for x in d["by_allocated_entity"]}
        self.assertEqual(alloc, {"ALPHA": "600.00", "BETA": "400.00"})

    def test_audit_is_append_only(self):
        a = AuditEvent.objects.create(action="x")
        a.action = "y"
        with self.assertRaises(ValidationError):
            a.save()
        with self.assertRaises(ValidationError):
            a.delete()

    def test_google_sheets_stub(self):
        self.login("finance")
        self.assertEqual(self.client.get("/api/integrations/google-sheets/").json()["status"], "pending")
        self.assertEqual(self.client.post("/api/integrations/google-sheets/").status_code, 501)


def make_xlsx(rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@override_settings(PRIVATE_MEDIA_ROOT="/tmp/expense-test-media")
class ImportExportTests(BaseAPITest):
    def upload(self, content, name="exp.xlsx", reprocess=False):
        data = {"file": SimpleUploadedFile(name, content)}
        if reprocess:
            data["reprocess"] = "1"
        r = self.client.post("/api/imports/", data)
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()

    def row(self, **kw):
        base = {"ExternalRecordID": "SRC-1", "ExpenseName": "Taxi", "ExpenseDate": "2026-09-10",
                "PayingEntityCode": "ALPHA", "CategoryCode": "FUEL", "GrossAmount": "350", "Notes": "imported"}
        base.update(kw)
        return base

    def test_only_finance_can_import(self):
        self.login("employee")
        r = self.client.post("/api/imports/", {"file": SimpleUploadedFile("a.xlsx", make_xlsx([self.row()]))})
        self.assertEqual(r.status_code, 403)

    def test_preview_commit_idempotent_and_update(self):
        self.login("finance")
        content = make_xlsx([
            self.row(),
            self.row(ExternalRecordID="SRC-2", ExpenseName="Hotel", CategoryCode="LODGING", TaxableValue="5000",
                     GSTType="CGST_SGST", GSTRate="12", GrossAmount="5600", InvoiceNumber="H-1",
                     InvoiceDate="2026-09-09", VendorName="Taj Hotels"),
            self.row(ExternalRecordID="SRC-3", PayingEntityCode="NOPE"),
            self.row(ExternalRecordID="SRC-4", ExpenseDate="05/06/2026"),
            self.row(ExternalRecordID="SRC-5", GrossAmount="-3"),
        ])
        run = self.upload(content)
        self.assertEqual(run["status"], SyncStatus.AWAITING_REVIEW)
        self.assertEqual(run["row_counts"]["NEW"], 2)
        self.assertEqual(run["row_counts"]["REJECTED"], 3)
        self.assertEqual(Expense.objects.count(), 0)  # preview writes nothing
        err = self.client.get(f"/api/imports/{run['id']}/errors.csv/")
        codes = {r["ErrorCode"] for r in csv.DictReader(io.StringIO(err.content.decode()))}
        self.assertEqual(codes, {"UNKNOWN_CODE", "BAD_DATE", "BAD_NUMBER"})
        done = self.client.post(f"/api/imports/{run['id']}/commit/").json()
        self.assertEqual(done["status"], SyncStatus.PARTIAL)
        self.assertEqual(Expense.objects.count(), 2)
        hotel = Expense.objects.get(external_record_id="SRC-2")
        self.assertEqual(hotel.cgst_amount, Decimal("300.00"))
        self.assertEqual(hotel.status, Status.DRAFT)
        self.assertEqual(hotel.source_type, SourceType.EXCEL)

        # same file again: skipped
        again = self.upload(content)
        self.assertEqual(again["status"], SyncStatus.SKIPPED)
        # reprocess: rows unchanged, nothing duplicated
        re = self.upload(content, reprocess=True)
        self.assertEqual(re["row_counts"]["UNCHANGED"], 2)
        self.client.post(f"/api/imports/{re['id']}/commit/")
        self.assertEqual(Expense.objects.count(), 2)

        # changed row updates the existing expense
        upd = self.upload(make_xlsx([self.row(GrossAmount="400", Notes="corrected")]))
        self.assertEqual(upd["row_counts"]["UPDATE"], 1)
        self.client.post(f"/api/imports/{upd['id']}/commit/")
        taxi = Expense.objects.get(external_record_id="SRC-1")
        self.assertEqual(taxi.gross_amount, Decimal("400.00"))
        self.assertEqual(taxi.notes, "corrected")
        self.assertEqual(Expense.objects.count(), 2)

    def test_locked_records_not_overwritten(self):
        self.login("finance")
        run = self.upload(make_xlsx([self.row()]))
        self.client.post(f"/api/imports/{run['id']}/commit/")
        exp = Expense.objects.get()
        exp.status = Status.FINANCE_VERIFIED
        exp.save()
        upd = self.upload(make_xlsx([self.row(GrossAmount="999")]))
        self.assertEqual(upd["row_counts"]["REJECTED"], 1)
        rows = self.client.get(f"/api/imports/{upd['id']}/rows/").json()
        self.assertEqual(rows[0]["errors"][0]["code"], "LOCKED")

    def test_heuristic_duplicate_flagged_not_discarded(self):
        self.login("finance")
        r1 = self.row(ExternalRecordID="", InvoiceNumber="inv-77", InvoiceDate="2026-09-10", VendorName="Uber")
        run = self.upload(make_xlsx([r1]))
        self.client.post(f"/api/imports/{run['id']}/commit/")
        run2 = self.upload(make_xlsx([{**r1, "InvoiceNumber": "INV 77", "Notes": "again"}]))
        self.assertEqual(run2["row_counts"]["DUPLICATE"], 1)
        self.client.post(f"/api/imports/{run2['id']}/commit/")
        self.assertEqual(Expense.objects.count(), 2)
        self.assertEqual(Expense.objects.filter(duplicate_suspect=True).count(), 1)

    def test_csv_import(self):
        self.login("finance")
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=HEADERS)
        w.writeheader()
        w.writerow(self.row())
        run = self.upload(buf.getvalue().encode(), name="e.csv")
        self.client.post(f"/api/imports/{run['id']}/commit/")
        self.assertEqual(Expense.objects.get().source_type, SourceType.CSV)

    def test_bad_file_fails_cleanly(self):
        self.login("finance")
        run = self.upload(b"not a workbook", name="x.xlsx")
        self.assertEqual(run["status"], SyncStatus.FAILED)
        run = self.upload(make_xlsx([{}], headers=["Foo", "Bar"]))
        self.assertEqual(run["status"], SyncStatus.FAILED)

    def test_exports_and_round_trip(self):
        self.login("finance")
        run = self.upload(make_xlsx([self.row(), self.row(ExternalRecordID="SRC-9", GrossAmount="50")]))
        self.client.post(f"/api/imports/{run['id']}/commit/")
        x = self.client.get("/api/exports/expenses.xlsx?status_scope=all")
        self.assertEqual(x.status_code, 200)
        wb = load_workbook(io.BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Expenses", "Report info", "By category"])
        self.assertEqual(wb["Expenses"].max_row, 3)
        c = self.client.get("/api/exports/expenses.csv?status_scope=all")
        rows = list(csv.DictReader(io.StringIO(c.content.decode("utf-8-sig"))))
        self.assertEqual(sum(Decimal(r["GrossAmount"]) for r in rows), Decimal("400.00"))
        d = self.client.get("/api/dashboard/?status_scope=all").json()
        self.assertEqual(d["kpis"]["gross"], "400.00")  # export reconciles with screen
        p = self.client.get("/api/exports/expenses.pdf?status_scope=all")
        self.assertTrue(p.content.startswith(b"%PDF"))
        # Edit the exported workbook and re-import: updates, no duplicates
        ws = wb["Expenses"]
        col = [c.value for c in ws[1]].index("Notes") + 1
        ws.cell(row=2, column=col, value="edited in Excel")
        buf = io.BytesIO()
        wb.save(buf)
        r = self.upload(buf.getvalue(), name="export.xlsx")
        self.assertEqual(r["row_counts"]["UPDATE"], 1)
        self.assertEqual(r["row_counts"]["UNCHANGED"], 1)
        self.client.post(f"/api/imports/{r['id']}/commit/")
        self.assertEqual(Expense.objects.count(), 2)
        self.assertTrue(Expense.objects.filter(notes="edited in Excel").exists())
        self.assertTrue(AuditEvent.objects.filter(action="export.xlsx").exists())
        t = self.client.get("/api/imports-template.xlsx")
        self.assertEqual(load_workbook(io.BytesIO(t.content))["Expenses"].cell(1, 1).value, "ExpenseRef")

    def test_export_respects_scope(self):
        self.login("finance")
        run = self.upload(make_xlsx([self.row()]))
        self.client.post(f"/api/imports/{run['id']}/commit/")
        self.login("employee2")
        rows = list(csv.DictReader(io.StringIO(
            self.client.get("/api/exports/expenses.csv?status_scope=all").content.decode("utf-8-sig"))))
        self.assertEqual(rows, [])


class SeedTests(TestCase):
    def test_seed_with_expenses_and_pages_render(self):
        call_command("seed_demo", "--count", "25", stdout=io.StringIO())
        self.assertEqual(Expense.objects.count(), 25)
        self.client.force_login(User.objects.get(username="finance"))
        self.assertEqual(self.client.get("/").status_code, 200)
        d = self.client.get("/api/dashboard/").json()
        self.assertGreater(Decimal(d["kpis"]["gross"]), 0)
        self.assertEqual(self.client.get("/api/reports/gst-summary/").status_code, 200)
        self.assertEqual(self.client.get("/api/audit/").status_code, 200)
